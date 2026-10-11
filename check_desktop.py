# -*- coding: utf-8 -*-
import io
from docx import Document

d = Document(r"C:\Users\ASUS\Desktop\桌面简历\技术策划简历.docx")
txt = "\n".join(p.text for p in d.paragraphs)
out = io.open(r"E:\0\desktop_check.txt", "w", encoding="utf-8")
out.write("has_refactor_title: %s\n" % ("像素回合制" in txt))
out.write("has_11_maps: %s\n" % ("11 张地图" in txt))
out.write("has_exe_delivery: %s\n" % ("单文件 EXE" in txt))
out.write("para_count: %d\n" % len(d.paragraphs))
out.write("---- project section ----\n")
for p in d.paragraphs[7:14]:
    out.write(p.text + "\n")
out.close()
print("done")
