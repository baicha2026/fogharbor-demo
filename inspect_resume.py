# -*- coding: utf-8 -*-
import io
from docx import Document
from docx.shared import Pt

path = r"C:\Users\ASUS\Desktop\桌面简历\技术策划简历.docx"
doc = Document(path)
out = io.open(r"E:\0\inspect_output.txt", "w", encoding="utf-8")
def w(s=""):
    out.write(s + "\n")

w("=== SECTIONS ===")
for i, sec in enumerate(doc.sections):
    w(f"Section {i}: page_w={sec.page_width}, page_h={sec.page_height}, "
      f"L={sec.left_margin}, R={sec.right_margin}, T={sec.top_margin}, B={sec.bottom_margin}")

w("\n=== PARAGRAPHS ===")
for idx, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else "None"
    align = str(p.alignment)
    txt = p.text
    runinfo = []
    for r in p.runs:
        fname = r.font.name
        fsize = r.font.size.pt if r.font.size else None
        bold = r.font.bold
        runinfo.append((r.text[:20], fname, fsize, bold))
    w(f"[{idx}] style={style} align={align} | {txt!r}")
    if runinfo:
        w(f"     runs={runinfo}")

w("\n=== TABLES ===")
for ti, t in enumerate(doc.tables):
    w(f"Table {ti}: rows={len(t.rows)}, cols={len(t.columns)}")
    for ri, row in enumerate(t.rows):
        cells = [c.text for c in row.cells]
        w(f"  row{ri}: {cells}")

out.close()
print("done")
