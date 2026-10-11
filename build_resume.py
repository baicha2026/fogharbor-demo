# -*- coding: utf-8 -*-
"""生成精简排版后的技术策划简历（1 页 A4，融合最新求职背景）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "微软雅黑"
DARK = "1F1F1F"
ACCENT = "1F4E79"
GRAY = "595959"

doc = Document()

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.0)
sec.bottom_margin = Cm(0.8)
sec.left_margin = Cm(1.5)
sec.right_margin = Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def set_run(run, size=10, bold=False, color=None, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_bottom_border(paragraph, color="BBBBBB", sz="6"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def para(align=None, before=0, after=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    return p


def heading(text, before=5):
    p = para(before=before, after=2)
    r = p.add_run(text)
    set_run(r, size=11, bold=True, color=ACCENT)
    add_bottom_border(p)
    return p


def bullet(bold_head=None, body=None, indent=0.35, after=1):
    p = para(after=after)
    pf = p.paragraph_format
    pf.left_indent = Cm(indent)
    pf.first_line_indent = Cm(-indent)
    r0 = p.add_run("•  ")
    set_run(r0, size=10, bold=False, color=DARK)
    if bold_head:
        rb = p.add_run(bold_head)
        set_run(rb, size=10, bold=True, color=DARK)
    if body:
        rn = p.add_run(body)
        set_run(rn, size=10, bold=False, color=DARK)
    return p


# ================= 头部 =================
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
r = p.add_run("王智涛")
set_run(r, size=16, bold=True, color=DARK)

p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
r = p.add_run("求职意向：技术策划 / 战斗策划 / 关卡策划 / 任务策划")
set_run(r, size=10.5, bold=True, color=ACCENT)

p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
r = p.add_run("电话 13019045796 ｜ 邮箱 baicha1913258369@163.com ｜ 意向城市：杭州 ｜ 招聘类型：秋招")
set_run(r, size=9.5, bold=False, color=GRAY)

p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
r = p.add_run("黑龙江工程学院 · 数据科学与大数据技术（本科）· 2027 届")
set_run(r, size=9.5, bold=False, color=GRAY)

# ================= 个人总结 =================
heading("个人总结")
bullet("技术策划复合能力：", "既能拆解战斗 / 关卡 / 任务系统，又能用 Python + Vibe Coding 把策划案快速做成可运行的数据沙盘。")
bullet("重度玩家：", "3000+ 小时游戏经历，覆盖 MOBA / FPS / 卡牌 / 肉鸽 / MMORPG / 开放世界，擅长从玩家与策划双视角拆解系统。")

# ================= 核心项目经历 =================
heading("核心项目经历")
p = para(after=0)
r = p.add_run("《雾港守灯人》｜2D 像素回合制冒险游戏 + AI 策划验证台")
set_run(r, size=10.5, bold=True, color=DARK)
p = para(after=1)
r = p.add_run("独立技术策划 / 系统原型开发　｜　端到端交付「任务生成 → 关卡探索 → 战斗结算 → 数据复盘」可运行闭环")
set_run(r, size=9.5, bold=False, color=GRAY)

bullet("任务策划｜完整流程与 AI 动态任务：",
       "搭建「主线 → 支线 → 商店 → 首领 → 结局」回合制流程；试炼场保留 AI 动态任务，Prompt 约束 + JSON 结构化输出 + Pydantic 校验，防止生成「超模任务」。")
bullet("关卡策划｜多路线与受击热力图：",
       "设计 11 张地图、「守灯 / 逐潮 / 破誓」3 条路线与 3 个结局，路线选择真实影响后续地图、敌人、技能书与奖励；用移动轨迹与受击热力图验证难度曲线。")
bullet("战斗策划｜双技能循环与 DPS 验证：",
       "设计「破阵重击（高伤长 CD）＋潮汐回响（低伤回蓝）」双技能循环与资源模型；用 DPS 曲线定位技能真空期，设计首领蓄力预告与防御减伤，验证心流节奏。")
bullet("技术策划｜端到端交付：",
       "Streamlit + FastAPI + Plotly / Pandas 实现前端与图表；接入 DeepSeek 异步复盘（≤200 字建议）与 ComfyUI 成就海报；打包 Windows 单文件 EXE 免安装运行，pytest + 浏览器自动化覆盖数值与交互。")

# ================= 深度游戏经历 =================
heading("深度游戏经历与拆解")
bullet("硬核动作 / 策略：",
       "《死亡细胞》300h · 5 细胞白王（伤害乘区 / 无敌帧 / 高风险高收益心流）；《杀戮尖塔》200h · 全角色 A20 碎心（数值螺旋 / Build 构筑）。")
bullet("MMORPG：",
       "《命运方舟》《诛仙世界》，深谙副本掉落、日常任务链与养成节奏，能拆解「中期资源溢出、日常枯燥」痛点并提出「动态难度 + 资源回收」方案。")
bullet("MOBA / FPS：",
       "《王者荣耀》（多赛季王者）《英雄联盟》《穿越火线》《三角洲行动》，精通技能 CD 循环、团战博弈、地图控制与射击手感反馈。")
bullet("开放世界 / 沙盒：",
       "《原神》；《我的世界》500h+，长期实践红石与命令方块，具备系统逻辑搭建能力。")

# ================= 专业技能 =================
heading("专业技能")
bullet("策划能力：", "系统架构、战斗数值建模、技能循环、关卡动线与心流、任务网搭建、竞品拆解。")
bullet("技术策划：", "Python、FastAPI、MySQL、Redis、Streamlit / Plotly / Pandas、Vibe Coding（Cursor / Codex）。")
bullet("AI 工作流：", "ComfyUI 节点工作流与 API 集成；LLM Agent Prompt 设计、JSON 输出约束与校验、Function Calling。")
bullet("工具软件：", "Excel（数值建模）、Visio / XMind（系统架构）、Axure / Figma（UI 原型）。")

# ================= 教育背景 =================
heading("教育背景")
bullet("黑龙江工程学院 · 数据科学与大数据技术 · 本科 · 2027 年毕业", after=1)
bullet("实习经历：专注游戏拆解与策划原型开发", after=1)

out_path = r"E:\0\技术策划简历.docx"
doc.save(out_path)
print("saved")
