# -*- coding: utf-8 -*-
import io
from docx import Document

doc = Document(r"E:\0\技术策划简历.docx")
out = io.open(r"E:\0\final_content.txt", "w", encoding="utf-8")
for idx, p in enumerate(doc.paragraphs):
    out.write(f"[{idx}] {p.text}\n")
out.close()

# 检查占位符与敏感词
full = "\n".join(p.text for p in doc.paragraphs)
issues = []
for kw in ["待补", "待填写", "待确认", "待定", "[待", "【待", "TODO", "xxx", "XX"]:
    if kw in full:
        issues.append(kw)
print("placeholder issues:", issues if issues else "NONE")
print("has email:", "baicha1913258369@163.com" in full)
print("has name:", "王智涛" in full)
print("paragraph count:", len(doc.paragraphs))
